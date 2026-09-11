"""Akt-sverka (statement) natijasini .xlsx faylga aylantirish uchun umumiy
yordamchi. `apps/reports/services/statement_service.py` qaytargan
dict'lardan (`opening_balance`, `rows`, `closing_balance`) foydalanadi."""
from openpyxl import Workbook
from openpyxl.styles import Font
from django.http import HttpResponse
from django.utils import timezone


def _write_sheet(ws, title, statement):
    ws.append([title])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Boshlang'ich saldo", float(statement['opening_balance'])])
    ws.append([])

    header = ['Sana', 'Operatsiya', 'Turkum', 'Kirim (+)', 'Chiqim (-)', 'Balans']
    ws.append(header)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for row in statement['rows']:
        ws.append([
            row['date'].strftime('%d.%m.%Y'),
            row['op'],
            row.get('category', ''),
            float(row['debit']),
            float(row['credit']),
            float(row['balance']),
        ])

    ws.append([])
    ws.append(['Yakuniy saldo', '', '', '', '', float(statement['closing_balance'])])
    ws[f'A{ws.max_row}'].font = Font(bold=True)

    for column_cells in ws.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 40)


def statement_to_response(*, filename, sections):
    """`sections` — [(sheet_title, statement_dict), ...]. Kamida bitta bo'lishi kerak."""
    wb = Workbook()
    for i, (title, statement) in enumerate(sections):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title[:31]
        _write_sheet(ws, title, statement)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def table_to_response(*, filename_prefix, headers, rows, sheet_title='Sheet1'):
    """Oddiy jadval (headers + rows) ni .xlsx faylga aylantiradi.

    Fayl nomiga yuklab olingan sana va soat qo'shiladi
    (masalan `mijozlar_2026-09-11_14-30-05.xlsx`), shu bilan qaysi vaqtda
    olingan hisobot ekani faylning o'zidan ko'rinib turadi.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)

    for column_cells in ws.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 40)

    timestamp = timezone.localtime().strftime('%Y-%m-%d_%H-%M-%S')
    filename = f'{filename_prefix}_{timestamp}.xlsx'

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
