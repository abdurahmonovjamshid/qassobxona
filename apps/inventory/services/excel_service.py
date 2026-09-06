"""Ombor qoldig'ini Excel (.xlsx) fayl orqali eksport/import qilish.

Foydalanuvchi butun omborni Excel'ga yuklab oladi, kg/dona/tannarx
ustunlarini to'g'rilab, qayta yuklaydi — tizim buni oddiy Inventarizatsiya
(`InventoryCount`) sifatida qayta ishlaydi, shu bilan tarix va validatsiya
saqlanib qoladi (`inventory_service.confirm_inventory_count`ga qarang).
"""
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from django.http import HttpResponse

from apps.inventory.services import inventory_service
from apps.products.models import Product

HEADER = ['ID', 'Kod', 'Nomi', 'Netto qoldiq (kg)', 'Dona', "Tannarx (so'm/kg)"]


def export_inventory_workbook() -> HttpResponse:
    """Faol mahsulotlarning joriy qoldig'i va o'rtacha tannarxini .xlsx
    faylga chiqaradi — shu fayl keyin to'g'irlanib qayta import qilinishi
    mumkin. `ID`, `Kod`, `Nomi` ustunlari faqat ma'lumot uchun — ularni
    o'zgartirmang, import ID bo'yicha moslashtiradi."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ombor'
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)

    for product in Product.objects.filter(active=True).select_related('category').order_by('name'):
        ws.append([
            product.id,
            product.code,
            product.name,
            float(inventory_service.get_stock(product)),
            inventory_service.get_stock_pieces(product),
            float(inventory_service.get_weighted_average_cost(product)),
        ])

    for column_cells in ws.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="ombor.xlsx"'
    wb.save(response)
    return response


def parse_inventory_workbook(file) -> tuple[list[dict], list[str]]:
    """Yuklangan .xlsx faylni o'qib, `{'product', 'kg', 'pieces', 'unit_cost'}`
    lug'atlari ro'yxatini qaytaradi. `unit_cost` — agar ustun bo'sh qoldirilgan
    bo'lsa `None` (tannarx o'zgarmaydi), aks holda shu qatorning tannarxi
    qayta belgilanadi. Xatolik topilsa, bo'sh ro'yxat va xatolar ro'yxati
    qaytariladi (hech narsa import qilinmaydi)."""
    errors = []
    try:
        wb = load_workbook(file, data_only=True)
    except Exception:
        return [], ["Fayl .xlsx formatida emas yoki buzilgan."]

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], ["Fayl bo'sh."]

    header = [str(c).strip() if c is not None else '' for c in rows[0][:len(HEADER)]]
    if header != HEADER:
        errors.append(
            f"Fayl sarlavhasi noto'g'ri. Kutilgan ustunlar: {', '.join(HEADER)}. "
            "Avval 'Excel shablonini yuklab olish' orqali joriy omborni eksport qiling."
        )
        return [], errors

    products_by_id = Product.objects.in_bulk()
    parsed = []
    for row_num, row in enumerate(rows[1:], start=2):
        if row is None or all(c is None or str(c).strip() == '' for c in row):
            continue

        raw_id, _code, name, raw_kg, raw_pieces, raw_cost = (list(row) + [None] * 6)[:6]

        try:
            product_id = int(raw_id)
        except (TypeError, ValueError):
            errors.append(f"{row_num}-qator: ID noto'g'ri yoki bo'sh.")
            continue
        product = products_by_id.get(product_id)
        if not product:
            errors.append(f"{row_num}-qator: ID={raw_id} bo'yicha mahsulot topilmadi ({name or ''}).")
            continue

        try:
            kg = Decimal(str(raw_kg)) if raw_kg is not None and str(raw_kg).strip() != '' else Decimal('0')
        except InvalidOperation:
            errors.append(f"{row_num}-qator ({product.name}): Netto (kg) noto'g'ri son.")
            continue
        if kg < 0:
            errors.append(f"{row_num}-qator ({product.name}): Netto (kg) manfiy bolmasligi kerak.")
            continue

        try:
            pieces = int(raw_pieces) if raw_pieces is not None and str(raw_pieces).strip() != '' else 0
        except (TypeError, ValueError):
            errors.append(f"{row_num}-qator ({product.name}): Dona noto'g'ri son.")
            continue
        if pieces < 0:
            errors.append(f"{row_num}-qator ({product.name}): Dona manfiy bolmasligi kerak.")
            continue

        unit_cost = None
        if raw_cost is not None and str(raw_cost).strip() != '':
            try:
                unit_cost = Decimal(str(raw_cost))
            except InvalidOperation:
                errors.append(f"{row_num}-qator ({product.name}): Tannarx noto'g'ri son.")
                continue
            if unit_cost < 0:
                errors.append(f"{row_num}-qator ({product.name}): Tannarx manfiy bolmasligi kerak.")
                continue

        parsed.append({'product': product, 'kg': kg, 'pieces': pieces, 'unit_cost': unit_cost})

    if errors:
        return [], errors
    if not parsed:
        errors.append("Faylda hech qanday qator topilmadi.")
        return [], errors
    return parsed, []
