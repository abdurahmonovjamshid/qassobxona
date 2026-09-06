from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.inventory.models import InventoryCount, InventoryCountItem
from apps.inventory.services import excel_service, inventory_service
from apps.products.models import Product, ProductCategory


def _make_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(excel_service.HEADER)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


class ConfirmInventoryCountCostTests(TestCase):
    """`confirm_inventory_count` — unit_cost berilganda tannarxni ham
    qayta belgilash (to'liq chiqim + qayta kirim) rejimini tekshiradi."""

    def setUp(self):
        category = ProductCategory.objects.create(name='Go‘sht', code='meat')
        self.product = Product.objects.create(name='Lahm', code='LAHM', category=category)
        inventory_service.stock_in(
            product=self.product, quantity=Decimal('100'), movement_type='PURCHASE',
            unit_cost=Decimal('50000'), date=timezone.localdate(),
        )

    def _confirm(self, counted_kg, counted_pieces, unit_cost):
        count = InventoryCount.objects.create(date=timezone.localdate())
        item = InventoryCountItem.objects.create(
            count=count, product=self.product, counted_kg=counted_kg,
            counted_pieces=counted_pieces, unit_cost=unit_cost,
        )
        inventory_service.confirm_inventory_count(count)
        return count, item

    def test_diff_only_when_unit_cost_blank_preserves_average(self):
        self._confirm(Decimal('95'), 0, None)
        self.assertEqual(inventory_service.get_stock(self.product), Decimal('95'))
        self.assertEqual(inventory_service.get_weighted_average_cost(self.product), Decimal('50000.00'))

    def test_unit_cost_provided_shifts_weighted_average_toward_new_cost(self):
        self._confirm(Decimal('95'), 2, Decimal('55000'))
        self.assertEqual(inventory_service.get_stock(self.product), Decimal('95'))
        self.assertEqual(inventory_service.get_stock_pieces(self.product), 2)

        expected_avg = (
            (Decimal('100') * Decimal('50000') + Decimal('95') * Decimal('55000'))
            / (Decimal('100') + Decimal('95'))
        ).quantize(Decimal('0.01'))
        self.assertEqual(inventory_service.get_weighted_average_cost(self.product), expected_avg)
        # Yangi tannarx eskisiga qaraganda o'rtachani yangi qiymat tomon suradi.
        self.assertGreater(expected_avg, Decimal('50000'))
        self.assertLess(expected_avg, Decimal('55000'))


class InventoryExcelRoundTripTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username='u1', password='pass12345')
        self.client.force_login(self.user)
        category = ProductCategory.objects.create(name='Go‘sht', code='meat')
        self.product = Product.objects.create(name='Lahm', code='LAHM', category=category)
        inventory_service.stock_in(
            product=self.product, quantity=Decimal('100'), movement_type='PURCHASE',
            unit_cost=Decimal('50000'), date=timezone.localdate(),
        )

    def test_export_contains_current_stock_row(self):
        response = self.client.get(reverse('inventory:export'))
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0], tuple(excel_service.HEADER))
        data_row = next(r for r in rows[1:] if r[0] == self.product.id)
        self.assertEqual(data_row[3], 100.0)
        self.assertEqual(data_row[5], 50000.0)

    def test_import_updates_stock_and_cost(self):
        content = _make_xlsx([[self.product.id, self.product.code, self.product.name, 95, 2, 55000]])
        upload = SimpleUploadedFile(
            'ombor.xlsx', content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response = self.client.post(reverse('inventory:import'), {'file': upload})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(inventory_service.get_stock(self.product), Decimal('95'))
        self.assertEqual(inventory_service.get_stock_pieces(self.product), 2)
        self.assertEqual(InventoryCount.objects.count(), 1)

    def test_import_leaves_cost_untouched_when_blank(self):
        content = _make_xlsx([[self.product.id, self.product.code, self.product.name, 80, 0, None]])
        upload = SimpleUploadedFile('ombor.xlsx', content)
        self.client.post(reverse('inventory:import'), {'file': upload})
        self.assertEqual(inventory_service.get_stock(self.product), Decimal('80'))
        self.assertEqual(inventory_service.get_weighted_average_cost(self.product), Decimal('50000.00'))

    def test_import_rejects_unknown_product_id(self):
        content = _make_xlsx([[999999, 'XXX', 'Yoq mahsulot', 10, 0, 1000]])
        upload = SimpleUploadedFile('ombor.xlsx', content)
        response = self.client.post(reverse('inventory:import'), {'file': upload}, follow=True)
        self.assertEqual(InventoryCount.objects.count(), 0)
        messages = list(response.context['messages'])
        self.assertTrue(any('topilmadi' in str(m) for m in messages))

    def test_import_rejects_negative_kg(self):
        content = _make_xlsx([[self.product.id, self.product.code, self.product.name, -5, 0, 1000]])
        upload = SimpleUploadedFile('ombor.xlsx', content)
        response = self.client.post(reverse('inventory:import'), {'file': upload}, follow=True)
        self.assertEqual(InventoryCount.objects.count(), 0)
        messages = list(response.context['messages'])
        self.assertTrue(any('manfiy' in str(m) for m in messages))

    def test_import_rejects_wrong_header(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['Wrong', 'Header'])
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        upload = SimpleUploadedFile('bad.xlsx', buf.read())
        response = self.client.post(reverse('inventory:import'), {'file': upload}, follow=True)
        self.assertEqual(InventoryCount.objects.count(), 0)
        messages = list(response.context['messages'])
        self.assertTrue(any('sarlavhasi' in str(m) for m in messages))
